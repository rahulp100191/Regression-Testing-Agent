import io
import json
import uuid
from pathlib import Path
from urllib.request import Request, urlopen
import openpyxl

root = Path(__file__).resolve().parents[2]
release = Path("C:/Users/HP/Downloads/WD release notes/What's_New_in_Workday (9).xlsx")
catalog = root / 'e2e list.xlsx'

def post(endpoint, release_bytes):
    boundary = uuid.uuid4().hex
    body = bytearray()
    for field, name, data in [('release_file','release.xlsx',release_bytes),('e2e_file','e2e.xlsx',catalog.read_bytes())]:
        body.extend(f'--{boundary}\r\nContent-Disposition: form-data; name="{field}"; filename="{name}"\r\nContent-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n'.encode())
        body.extend(data)
        body.extend(b'\r\n')
    body.extend(f'--{boundary}--\r\n'.encode())
    request = Request('http://127.0.0.1:8001'+endpoint,data=body,headers={'Content-Type':'multipart/form-data; boundary='+boundary,'Origin':'http://localhost:3000'},method='POST')
    with urlopen(request,timeout=180) as response:
        assert response.headers.get('Access-Control-Allow-Origin')=='http://localhost:3000'
        return json.load(response)

preview = post('/api/preview',release.read_bytes())
assert preview['release']['row_count']==65
assert preview['e2e']['row_count']==42
source = openpyxl.load_workbook(release,data_only=True).active
wb = openpyxl.Workbook()
wb.active.append([c.value for c in source[1]])
wb.active.append([c.value for c in source[65]])
buffer = io.BytesIO()
wb.save(buffer)
analysis = post('/api/analyze',buffer.getvalue())
assert analysis['release_count']==1
result = analysis['results'][0]
assert result['release']['title']=='Workday Time Kiosk Android Support'
names = {r[2] for r in openpyxl.load_workbook(catalog,data_only=True).active.iter_rows(min_row=2,values_only=True) if r[2]}
assert result['e2e_name'] is None or result['e2e_name'] in names
summary = {'ui_http_status':200,'frontend_api_base':'http://localhost:8001','browser_visual_verification':'unavailable: Windows Application Control blocked agent-browser, and browser connector has no available browser', 'preview':preview,'single_note_analysis':analysis}
(Path(__file__).parent/'verification.json').write_text(json.dumps(summary,indent=2,ensure_ascii=False)+'\n',encoding='utf-8')
print(json.dumps({'preview_release_rows':65,'preview_catalog_rows':42,'analyzed_title':result['release']['title'],'decision':result['decision'],'e2e':result['e2e_name'],'model_error':result.get('model_error'),'cors_origin_verified':True},indent=2))
raise SystemExit(2 if result.get('model_error') else 0)
