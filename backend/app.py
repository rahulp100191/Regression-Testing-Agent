from __future__ import annotations

import io
import json
import os
import re
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any
from uuid import UUID

from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from openpyxl import load_workbook
from dotenv import load_dotenv
from providers import generate, model_name, provider_name

load_dotenv(Path(__file__).resolve().parents[1] / ".env")
from telemetry import initialize, stage, event, RequestTracingMiddleware, trace_id, failure
initialize('regression-api')

RELEASE_COLUMNS = {
    "functional_area": "Functional Area (Domain)", "product_line": "Product Line", "product_area": "Product Area",
    "using_workday": "Using Workday", "title": "Title", "description": "Description",
    "business_benefits": "Business Benefits", "changes": "Changes", "impact": "Impact",
    "release_number": "Release Number", "release_note_url": "Release Note URL",
}
E2E_COLUMNS = {"item": "Item", "work_stream": "Work-Stream", "e2e": "E2E"}
RELEASE_2025_COLUMNS = {
    'functional_area': 'Functional Area(s)', 'title': 'New Functionality Title',
    'description': 'Feature Description', 'changes': 'New Functionality',
    'impact': 'Training & Testing Impact', 'release_number': 'Workday Release',
    'release_note_url': 'Community Post',
}
STOP_WORDS = {"the", "and", "for", "with", "from", "that", "this", "into", "your", "are", "now", "can", "you", "workday", "users", "user", "new", "using", "use", "when", "more", "only", "all", "not", "will", "their", "they", "has", "have", "been", "than"}
WORKSTREAM_HINTS = {
    "payroll": {"payroll", "pay", "payment", "tax", "wage", "settlement", "earning", "deduction"},
    "hcm core": {"hcm", "human", "capital", "staffing", "job", "worker", "employee", "organization", "hire", "termination"},
    "compensation": {"compensation", "salary", "bonus", "reward", "merit", "pay"},
    "advanced compensation": {"compensation", "salary", "bonus", "reward", "merit"},
    "absence": {"absence", "leave", "time off", "return from leave"},
    "time tracking": {"time", "hours", "reported", "elapsed", "exception"},
    "benefits": {"benefit", "enrollment", "coverage", "dependent"},
    "recruiting": {"recruit", "candidate", "job requisition", "interview", "offer", "hire"},
    "talent and performance": {"talent", "performance", "succession", "objective", "development", "disciplinary"},
    "learning": {"learning", "course", "program", "lesson", "topic"},
    "expenses": {"expense", "credit card", "expense report", "reimbursement"},
}


def clean(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value)).strip() if value is not None else ""


def normalize_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", clean(value).lower())


def tokens(value: str) -> set[str]:
    return {word for word in re.findall(r"[a-z0-9]+", value.lower()) if len(word) > 2 and word not in STOP_WORDS}


def table_from_upload(contents: bytes, kind: str) -> dict[str, Any]:
    import zipfile
    if len(contents) > int(os.getenv('MAX_UPLOAD_BYTES','10485760')):
        raise HTTPException(413, 'Workbook exceeds upload size limit')
    try:
        with zipfile.ZipFile(io.BytesIO(contents)) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > int(os.getenv('MAX_WORKBOOK_EXPANDED_BYTES','104857600')):
                raise HTTPException(413, 'Expanded workbook exceeds size limit')
    except zipfile.BadZipFile:
        raise HTTPException(400, 'Invalid XLSX archive') from None
    try:
        workbook = load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Could not read workbook: {exc}") from exc
    sheet = next((item for item in workbook.worksheets if item.max_row and item.max_column), None)
    if sheet is None:
        raise HTTPException(status_code=400, detail=f"{kind.title()} workbook has no non-empty worksheets")
    rows = list(sheet.iter_rows(values_only=True))
    headers = [clean(value) for value in rows[0]] if rows else []
    header_index = {normalize_header(header): index for index, header in enumerate(headers) if header}
    required = RELEASE_COLUMNS if kind == "release" else E2E_COLUMNS
    if kind == 'release' and normalize_header('New Functionality Title') in header_index:
        required = RELEASE_2025_COLUMNS
    missing = [label for label in required.values() if normalize_header(label) not in header_index]
    if missing:
        raise HTTPException(status_code=400, detail=f"{kind.title()} workbook is missing columns: {', '.join(missing)}")
    data = []
    if len(rows) > int(os.getenv('MAX_WORKBOOK_ROWS','20000')):
        raise HTTPException(413, 'Workbook exceeds row limit')
    for source_row, row in enumerate(rows[1:], start=2):
        if not any(clean(value) for value in row):
            continue
        record = {key: clean(row[header_index[normalize_header(label)]]) for key, label in required.items()}
        if kind == 'release':
            record = {**{key: '' for key in RELEASE_COLUMNS}, **record}
        record["source_row"] = source_row
        data.append(record)
    return {"sheet": sheet.title, "sheets": workbook.sheetnames, "headers": headers, "rows": data, "row_count": len(data)}


def release_text(row: dict[str, Any]) -> str:
    return " ".join(row.get(key, "") for key in RELEASE_COLUMNS if key not in {"release_note_url", "release_number"})


def shortlist(release: dict[str, Any], e2e_rows: list[dict[str, Any]], limit: int = 8) -> list[dict[str, Any]]:
    source = release_text(release)
    source_tokens = tokens(source)
    functional = clean(release.get("functional_area", "")).lower()
    product = clean(release.get("product_area", "")).lower()
    scored = []
    for candidate in e2e_rows:
        name = candidate["e2e"]
        candidate_text = f"{candidate.get('work_stream', '')} {name}"
        overlap = len(source_tokens & tokens(candidate_text)) / max(1, len(source_tokens | tokens(candidate_text)))
        stream = clean(candidate.get("work_stream", "")).lower()
        hint_words = WORKSTREAM_HINTS.get(stream, set())
        hint_bonus = 0.22 if any(hint in source.lower() for hint in hint_words) else 0.0
        if stream in functional or stream in product:
            hint_bonus += 0.28
        fuzzy = SequenceMatcher(None, clean(release.get("title", "")).lower(), name.lower()).ratio()
        scored.append((overlap * 0.55 + hint_bonus + fuzzy * 0.2, candidate))
    scored.sort(key=lambda item: item[0], reverse=True)
    return [{**candidate, "shortlist_score": round(score, 4)} for score, candidate in scored[:limit]]


def bedrock_model_id() -> str:
    return os.getenv("BEDROCK_MODEL_ID", "amazon.nova-lite-v1:0").strip()


def review_threshold() -> int:
    try:
        return int(os.getenv("REVIEW_THRESHOLD", "70"))
    except ValueError:
        return 70


def fallback_result(reason: str, candidates: list[dict[str, Any]]) -> dict[str, Any]:
    return {"decision": "unable_to_identify", "business_process_name": None, "e2e_name": None, "confidence": 0, "confidence_band": "low", "reasoning": reason, "evidence": [], "review_required": True, "shortlisted_candidates": [c["e2e"] for c in candidates]}


def extract_json(text: str) -> dict[str, Any]:
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text.strip(), flags=re.IGNORECASE | re.DOTALL).strip()
    start, end = text.find("{"), text.rfind("}")
    if start < 0 or end <= start:
        raise ValueError("Model response did not contain a JSON object")
    value = json.loads(text[start:end + 1])
    if not isinstance(value, dict):
        raise ValueError("Model response was not an object")
    return value


def validate_result(value: dict[str, Any], candidates: list[dict[str, Any]] | list[str], source_text: str = "") -> dict[str, Any]:
    candidate_names = [candidate["e2e"] if isinstance(candidate, dict) else candidate for candidate in candidates]
    raw_decision = value.get('decision')
    decision = raw_decision if isinstance(raw_decision, str) and raw_decision in {"match", "unable_to_identify", "no_matching_e2e"} else "unable_to_identify"
    selected = value.get("e2e_name") if value.get("e2e_name") in candidate_names else None
    if decision == "match" and not selected:
        decision = "unable_to_identify"
    if decision != 'match':
        selected = None
    try:
        confidence = max(0, min(100, int(float(value.get("confidence", 0)))))
    except (ValueError, TypeError, OverflowError):
        confidence = 0
    raw_evidence = value.get("evidence") or []
    if isinstance(raw_evidence, str):
        raw_evidence = [raw_evidence]
    if not isinstance(raw_evidence, list):
        raw_evidence = []
    evidence = [clean(item) for item in raw_evidence if clean(item)][:5]
    evidence_grounding = round(100 * sum(1 for item in evidence if item.lower() in source_text.lower()) / max(1, len(evidence))) if source_text else 0
    ranked = [candidate for candidate in candidates if isinstance(candidate, dict)]
    selected_candidate = next((candidate for candidate in ranked if candidate.get("e2e") == selected), None)
    rank = next((index + 1 for index, candidate in enumerate(ranked) if candidate.get("e2e") == selected), None)
    top_score = float(ranked[0].get("shortlist_score", 0)) if ranked else 0
    selected_score = float(selected_candidate.get("shortlist_score", 0)) if selected_candidate else 0
    second_score = float(ranked[1].get("shortlist_score", 0)) if len(ranked) > 1 else 0
    lexical_relevance = round(min(100, max(0, selected_score / 1.05 * 100)))
    selection_agreement = 100 if selected and rank == 1 else 60 if selected else 0
    model_confidence = confidence
    confidence = round(0.35 * model_confidence + 0.40 * lexical_relevance + 0.15 * selection_agreement + 0.10 * evidence_grounding)
    if decision != "match":
        confidence = min(confidence, 49)
    band = "high" if confidence >= 80 else "medium" if confidence >= 60 else "low"
    review_required = True if decision != "match" else bool(value.get("review_required", False)) or confidence < review_threshold() or evidence_grounding < 20 or rank != 1
    return {"decision": decision, "business_process_name": clean(value.get("business_process_name")) or None, "e2e_name": selected, "confidence": confidence, "confidence_band": band, "reasoning": clean(value.get("reasoning")) or "The model did not provide sufficient grounded reasoning.", "evidence": evidence, "review_required": review_required, "evaluation": {"model_confidence": model_confidence, "lexical_relevance": lexical_relevance, "candidate_rank": rank, "candidate_margin": round(max(0, top_score - second_score), 4), "evidence_grounding": evidence_grounding, "selection_agreement": selection_agreement}}


def build_prompt(release: dict[str, Any], candidates: list[dict[str, Any]]) -> str:
    from string import Template
    template = (Path(__file__).resolve().parents[1] / "prompts/regression/v1/prompt.txt").read_text(encoding="utf-8")
    return Template(template).substitute(release=json.dumps(release, ensure_ascii=False), candidates=json.dumps([c["e2e"] for c in candidates], ensure_ascii=False))


async def ask_agent(release: dict[str, Any], candidates: list[dict[str, Any]], provider: str | None = None) -> dict[str, Any]:
    return await guarded_prediction(release, candidates, provider)


async def guarded_prediction(release, candidates, provider=None, prompt_template=None):
    from providers import usage_capture
    from safety import screen, guard_output
    usage = {}
    input_report = {'action':'block', 'labels':['screening_not_completed']}
    token = usage_capture.set(usage)
    try:
        with stage('safety.input'):
            input_report = await screen({'release':release, 'candidates':candidates})
            event('safety.classified', action=input_report['action'], labels=input_report['labels'])
        if input_report['action'] == 'block':
            result = fallback_result('Input blocked by safety screening; review the source document.', candidates)
            result.update(raw_prediction=None, model_error='input_guardrail_blocked', model_usage=usage, safety={'input':input_report})
            return result
        with stage('prompt.render'):
            if prompt_template is None:
                prompt = build_prompt(release,candidates)
            else:
                from string import Template
                prompt = Template(prompt_template).substitute(release=json.dumps(release,ensure_ascii=False), candidates=json.dumps([c['e2e'] for c in candidates],ensure_ascii=False))
            event('prompt.rendered',prompt_version='regression-v1',characters=len(prompt))
        text = await generate(prompt, provider)
        with stage('response.validate'):
            raw = extract_json(text)
            output_report = await guard_output(raw, candidates, clean(release_text(release)))
            result = validate_result(raw, candidates, clean(release_text(release))) if output_report['action'] == 'allow' else fallback_result('Output blocked by safety or grounding validation.', candidates)
            result['safety'] = {'input':input_report, 'output':output_report}
            event('response.validated',decision=result['decision'],review_required=result['review_required'])
        # Preserve pre-guardrail fields for honest hallucination and calibration evaluation.
        result['raw_prediction'] = raw
        result['model_error'] = None
        result['model_usage'] = usage
        return result
    except Exception as exc:
        reason = str(exc) if isinstance(exc, (ValueError, RuntimeError)) else 'Model request failed; check provider configuration.'
        result = fallback_result(reason, candidates)
        failure(exc)
        result.update({'raw_prediction': None, 'model_error': reason, 'model_usage':usage})
        result['safety'] = {'input':input_report}
        return result
    finally:
        usage_capture.reset(token)


async def ask_bedrock(release: dict[str, Any], candidates: list[dict[str, Any]]) -> dict[str, Any]:
    return await ask_agent(release, candidates, provider='bedrock')


app = FastAPI(title="Workday Release Regression Agent API", version="1.0.0")
app.add_middleware(CORSMiddleware, allow_origins=list({os.getenv("FRONTEND_ORIGIN", "http://localhost:3000"),"http://localhost:3000","http://127.0.0.1:3000"}), allow_credentials=False, allow_methods=["*"], allow_headers=["*"], expose_headers=['X-Trace-ID','X-Request-ID'])
app.add_middleware(RequestTracingMiddleware)


@app.get("/api/health")
async def health():
    return {"status": "ok", "provider": provider_name(), "model": model_name(), "bedrock_region": os.getenv("AWS_REGION", "ap-southeast-2"), "bedrock_model": bedrock_model_id(), "review_threshold": review_threshold()}

@app.get('/api/diagnostics')
def diagnostics():
    from providers import model_routes
    from eval_store import database
    try:
        with database() as db:
            db.execute('SELECT 1')
        storage = 'connected'
    except Exception as exc:
        failure(exc,component='postgres')
        storage = 'unavailable'
    routes = model_routes()
    local_models = []
    local_state = 'not_configured'
    if any(route['provider']=='ollama' for route in routes):
        try:
            from urllib.request import urlopen
            with stage('ollama.catalog'):
                with urlopen(os.getenv('OLLAMA_BASE_URL','http://127.0.0.1:11434').rstrip('/')+'/api/tags',timeout=3) as response:
                    local_models = [{'model':m['name'],'digest':m.get('digest')} for m in json.load(response).get('models',[])]
                available = {m['model'] for m in local_models}
                expected = [r['model'] for r in routes if r['provider']=='ollama']
                local_state = 'ready' if all(m in available for m in expected) else 'partially_ready' if any(m in available for m in expected) else 'weights_downloading_or_missing'
                for route in routes:
                    if route['provider']=='ollama':
                        route['readiness'] = 'ready' if route['model'] in available else 'downloading_or_missing'
                event('ollama.catalog_read',state=local_state,model_count=len(local_models))
        except Exception as exc:
            failure(exc,component='ollama')
            local_state = 'unavailable'
    return {'service':'regression-api','storage':storage,'model_routes':routes,'local_models':local_models,'local_model_state':local_state,
            'trace_id':trace_id(),'dashboard_url':os.getenv('JAEGER_PUBLIC_URL','http://127.0.0.1:16686')}


@app.post("/api/preview")
async def preview(release_file: UploadFile = File(...), e2e_file: UploadFile = File(...)):
    if not (release_file.filename or "").lower().endswith(".xlsx") or not (e2e_file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Both files must be .xlsx workbooks")
    with stage('sources.parse'):
        release = table_from_upload(await release_file.read(int(os.getenv('MAX_UPLOAD_BYTES','10485760'))+1), "release")
        e2e = table_from_upload(await e2e_file.read(int(os.getenv('MAX_UPLOAD_BYTES','10485760'))+1), "e2e")
    return {"release": {key: release[key] for key in ["sheet", "sheets", "row_count", "headers"]}, "e2e": {key: e2e[key] for key in ["sheet", "sheets", "row_count", "headers"]}}


@app.get('/api/analyses')
def saved_analyses():
    from eval_store import database
    with database() as db:
        return db.execute("""SELECT run_id,count(*) AS saved_rows,max(created_at) AS last_saved,
          max((versions->>'total_release_count')::integer) AS total_rows
          FROM analysis_rows GROUP BY run_id ORDER BY last_saved DESC LIMIT 100""").fetchall()


@app.get('/api/analyses/{run_id}')
def saved_analysis(run_id: UUID):
    from eval_store import database
    from analysis_batch import public_prediction
    with database() as db:
        rows=db.execute('SELECT * FROM analysis_rows WHERE run_id=%s ORDER BY source_row',(run_id,)).fetchall()
    if not rows:
        raise HTTPException(404,'Saved analysis not found')
    return {'run_id':str(run_id), 'catalog_names':[c['e2e'] for c in rows[0]['catalog']],
            'results':[public_prediction(r) for r in rows], 'saved_rows':len(rows),
            'total_rows':max((r['versions'].get('total_release_count',0) for r in rows),default=0) or None}


@app.post("/api/analyze")
async def analyze(release_file: UploadFile = File(...), e2e_file: UploadFile = File(...), resume_run_id: UUID | None = Form(default=None), analysis_run_id: UUID | None = Form(default=None)):
    if not (release_file.filename or "").lower().endswith(".xlsx") or not (e2e_file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Both files must be .xlsx workbooks")
    with stage('sources.parse'):
        release = table_from_upload(await release_file.read(int(os.getenv('MAX_UPLOAD_BYTES','10485760'))+1), "release")
        e2e = table_from_upload(await e2e_file.read(int(os.getenv('MAX_UPLOAD_BYTES','10485760'))+1), "e2e")
    from eval_store import database, encode
    from uuid import uuid4
    import hashlib
    from safety import screen, VERSION
    from eval_api import implementation_version
    from analysis_batch import completed_batches, settings, PROMPT, public_prediction
    catalog_report = await screen(e2e['rows'])
    if catalog_report['action'] == 'block':
        raise HTTPException(422, detail={'message':'Test catalog blocked by safety screening', 'safety':catalog_report})
    run_id = resume_run_id or analysis_run_id or uuid4()
    config=settings()
    dataset_hash=hashlib.sha256(json.dumps({'release':release['rows'],'catalog':e2e['rows']},sort_keys=True,ensure_ascii=False).encode()).hexdigest()
    versions = {'implementation':implementation_version(), 'safety':VERSION, 'provider':provider_name(), 'model':model_name(),
                'prompt_version':'regression-v2-batch', 'prompt_sha256':hashlib.sha256(PROMPT.read_bytes()).hexdigest(),
                'batch_settings':config, 'dataset_sha256':dataset_hash, 'total_release_count':len(release['rows'])}
    versions['gemma_thinking_level']=os.getenv('GEMMA_THINKING_LEVEL','minimal')
    try:
        with database() as db:
            db.execute('SELECT 1')
    except Exception:
        raise HTTPException(503, 'Analysis storage unavailable; start PostgreSQL before analyzing') from None
    # A lease prevents concurrent resume requests from generating/persisting the same rows.
    with database() as lease:
        if not lease.execute('SELECT pg_try_advisory_xact_lock(hashtextextended(%s,0)) AS acquired',(str(run_id),)).fetchone()['acquired']:
            raise HTTPException(409,'This analysis is already running')
        saved=lease.execute('SELECT * FROM analysis_rows WHERE run_id=%s ORDER BY source_row',(run_id,)).fetchall()
        if saved and not resume_run_id:
            raise HTTPException(409,'Analysis ID already exists; use resume_run_id to reuse saved rows')
        if resume_run_id and not saved:
            raise HTTPException(404,'Saved analysis not found')
        source={row['source_row']:row for row in release['rows']}
        for record in saved:
            if record['catalog'] != e2e['rows'] or record['release'] != source.get(record['source_row']) or record['versions'].get('dataset_sha256',dataset_hash) != dataset_hash:
                raise HTTPException(409,'Uploaded workbooks differ from the saved analysis; start a new run')
        reused={record['source_row'] for record in saved}
        remaining=[row for row in release['rows'] if row['source_row'] not in reused]
        results=[public_prediction(record) for record in saved]
        event('analysis.plan',run_id=str(run_id),total_rows=len(source),reused_rows=len(saved),pending_rows=len(remaining),**config)
        async for items,predictions in completed_batches(remaining,e2e['rows'],config):
            with database() as db:
                for row,candidates in items:
                    result=predictions[row['source_row']]
                    result.update(source_row=row['source_row'],release=row,shortlisted_candidates=[c['e2e'] for c in candidates])
                    row_id=uuid4()
                    db.execute('INSERT INTO analysis_rows(row_id,run_id,source_row,release,catalog,prediction,versions) VALUES (%s,%s,%s,%s,%s,%s,%s)',
                               (row_id,run_id,row['source_row'],encode(row),encode(e2e['rows']),encode(result),encode(versions)))
                    results.append(public_prediction({'row_id':row_id,'prediction':result}))
            event('analysis.progress',run_id=str(run_id),completed_rows=len(results),total_rows=len(source))
    results.sort(key=lambda result:result['source_row'])
    return {'run_id':str(run_id),'batch_settings':config,'reused_rows':len(saved), 'catalog_names':[c['e2e'] for c in e2e['rows']],
            'release_count':len(release['rows']),'e2e_count':len(e2e['rows']),'review_threshold':review_threshold(),'results':results}


# Imported after app initialization to avoid evaluator/app circular imports.
from eval_api import router as eval_router
app.include_router(eval_router)
from feedback_api import router as feedback_router
app.include_router(feedback_router)
